import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- Color Palette ---
NAVY = RGBColor(0x1F, 0x4E, 0x79)       # Primary Accent (#1F4E79)
TEAL = RGBColor(0x2E, 0x75, 0xB6)       # Secondary Accent (#2E75B6)
CHARCOAL = RGBColor(0x26, 0x26, 0x26)   # Body text (#262626)
MUTED = RGBColor(0x59, 0x59, 0x59)      # Subtitles & footnotes (#595959)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)      # White
SUCCESS_GREEN = RGBColor(0x28, 0xA7, 0x45)
WARNING_AMBER = RGBColor(0xD9, 0x77, 0x06)

HEX_NAVY = "1F4E79"
HEX_LIGHT_BG = "F2F5F9"
HEX_BORDER = "D3D8E0"
HEX_CALLOUT_BG = "F0F4F8"
HEX_CALLOUT_BORDER = "1F4E79"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, border_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{border_color}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    run = h.runs[0]
    run.font.name = "Calibri"
    if level == 1:
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = NAVY
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = TEAL
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = CHARCOAL
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_p(doc, text, bold_prefix=None, space_after=5, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Calibri"
        r_bold.font.size = Pt(10.5)
        r_bold.font.bold = True
        r_bold.font.color.rgb = CHARCOAL
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = CHARCOAL
    return p

def add_bullet(doc, text, bold_prefix=None, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = "Calibri"
        r_bold.font.size = Pt(10.5)
        r_bold.font.bold = True
        r_bold.font.color.rgb = CHARCOAL
    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = CHARCOAL
    return p

def add_callout(doc, text, title="ARCHITECTURAL NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_t = p.add_run(f"[{title}] ")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(10)
    r_t.font.bold = True
    r_t.font.color.rgb = NAVY
    
    r_b = p.add_run(text)
    r_b.font.name = "Calibri"
    r_b.font.size = Pt(10)
    r_b.font.color.rgb = CHARCOAL
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    
    # Format Header
    hdr_row = tbl.rows[0]
    hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(heading)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = WHITE
        
    # Format Data Rows
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        row_bg = HEX_LIGHT_BG if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            if row_bg != "FFFFFF":
                set_cell_background(cell, row_bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(val))
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.color.rgb = CHARCOAL
            
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(6)
    return tbl

print("Base helper functions defined successfully")
